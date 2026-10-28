import os, sys, re, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

# 1. Colors
C_NAVY = RGBColor(26, 54, 93)      # #1A365D Primary title
C_BLUE = RGBColor(37, 99, 235)     # #2563EB Learner
C_PURPLE = RGBColor(124, 58, 237)  # #7C3AED Manager
C_ORANGE = RGBColor(234, 88, 12)   # #EA580C Creator
C_GREEN = RGBColor(22, 163, 74)    # #16A34A Moderator
C_GRAY = RGBColor(75, 85, 99)      # #4B5563 Operations
C_RED = RGBColor(153, 27, 27)      # #991B1B Administrator
C_DARK = RGBColor(15, 23, 42)      # #0F172A Text main
C_MUTED = RGBColor(100, 116, 139)  # #64748B Secondary text
C_BG_CARD = RGBColor(248, 250, 252)# #F8FAFC Card background
C_BORDER = RGBColor(226, 232, 240) # #E2E8F0 Border
C_WHITE = RGBColor(255, 255, 255)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

def create_deck():
    prs = Presentation('docs/scratch/clean_template.pptx')
    
    # Remove existing slides from template
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]
        
    blank_layout = prs.slide_layouts[6] # BLANK layout
    
    def add_header(slide, cat_text, title_text, subtitle_text=""):
        # Header text frame
        txBox = slide.shapes.add_textbox(Inches(1.8), Inches(0.2), Inches(11.2), Inches(1.1))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # Category
        p0 = tf.paragraphs[0]
        p0.text = cat_text.upper()
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = C_BLUE
        p0.space_after = Pt(2)
        
        # Title
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        
        # Subtitle
        if subtitle_text:
            p1.space_after = Pt(2)
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = FONT_BODY
            p2.font.size = Pt(11)
            p2.font.color.rgb = C_MUTED
            
    def add_card(slide, left, top, width, height, bg_color=C_BG_CARD, border_color=C_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def add_footer(slide, text=""):
        if text:
            txBox = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.3))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = text
            p.font.name = FONT_BODY
            p.font.size = Pt(9)
            p.font.color.rgb = C_MUTED

    # ==================== SLIDE 1: COVER ====================
    s1 = prs.slides.add_slide(blank_layout)
    # Left Hero Box
    tbox = s1.shapes.add_textbox(Inches(1.8), Inches(1.5), Inches(6.5), Inches(4.5))
    tf = tbox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "LINGORA"
    p.font.name = FONT_HEADING
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_after = Pt(6)
    
    p = tf.add_paragraph()
    p.text = "Phân tích yêu cầu nền tảng học ngoại ngữ trực tuyến"
    p.font.name = FONT_HEADING
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p.space_after = Pt(10)
    
    p = tf.add_paragraph()
    p.text = "Từ nhu cầu người học đến yêu cầu nghiệp vụ và quy trình vận hành"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = C_MUTED
    p.space_after = Pt(24)
    
    p = tf.add_paragraph()
    p.text = "Học phần: Dự án 1 — Hệ thống Thông tin\nNhóm thực hiện · 09/2026"
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = C_DARK
    
    # Right Phone Mockup Card
    add_card(s1, Inches(8.5), Inches(1.3), Inches(4.0), Inches(5.0), bg_color=C_WHITE, border_color=C_BLUE)
    tbox_m = s1.shapes.add_textbox(Inches(8.8), Inches(1.6), Inches(3.4), Inches(4.4))
    tf_m = tbox_m.text_frame
    tf_m.word_wrap = True
    
    pm = tf_m.paragraphs[0]
    pm.text = "📱 LINGORA MOBILE LEARNING"
    pm.font.name = FONT_HEADING
    pm.font.size = Pt(10)
    pm.font.bold = True
    pm.font.color.rgb = C_BLUE
    pm.space_after = Pt(14)
    
    pm = tf_m.add_paragraph()
    pm.text = "Good morning"
    pm.font.name = FONT_HEADING
    pm.font.size = Pt(24)
    pm.font.bold = True
    pm.font.color.rgb = C_NAVY
    
    pm = tf_m.add_paragraph()
    pm.text = "/ɡʊd ˈmɔːrnɪŋ/ • [🔊 Nghe audio]"
    pm.font.name = FONT_BODY
    pm.font.size = Pt(11)
    pm.font.color.rgb = C_BLUE
    pm.space_after = Pt(4)
    
    pm = tf_m.add_paragraph()
    pm.text = "Chào buổi sáng"
    pm.font.name = FONT_BODY
    pm.font.size = Pt(13)
    pm.font.color.rgb = C_DARK
    pm.space_after = Pt(16)
    
    pm = tf_m.add_paragraph()
    pm.text = "💡 6 Vai trò hệ thống phối hợp:"
    pm.font.name = FONT_HEADING
    pm.font.size = Pt(10)
    pm.font.bold = True
    pm.font.color.rgb = C_NAVY
    pm.space_after = Pt(4)
    
    pm = tf_m.add_paragraph()
    pm.text = "• Learner (Học & Ôn tập ngắt quãng)\n• Organizational Manager (Khung & Xuất bản)\n• Lesson Creator (Soạn bài & Tải audio chuẩn)\n• Moderator (Kiểm duyệt độc lập 6 tiêu chí)\n• Operations Staff (Hỗ trợ người dùng)\n• Administrator (Phân quyền & Audit Log)"
    pm.font.name = FONT_BODY
    pm.font.size = Pt(9.5)
    pm.font.color.rgb = C_MUTED

    # ==================== SLIDE 2: AGENDA ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Mục lục chương trình", "Từ vấn đề thực tế đến yêu cầu hệ thống", "Lộ trình 5 phần logic dẫn dắt từ nỗi đau người học đến mô hình phần mềm")
    
    parts = [
        ("01 — BÀI TOÁN", "Slide 3", "Vấn đề hay quên, gián đoạn của người học và bài toán quản trị học liệu phía sau nền tảng.", C_BLUE),
        ("02 — DỮ LIỆU & THỰC CHỨNG", "Slide 4–10", "Khảo sát N=15, N=35 · Phỏng vấn định tính N=3 · Phỏng vấn chuyên môn 7 Stakeholders · Benchmark · Legal Review.", C_PURPLE),
        ("03 — ĐẶC TẢ YÊU CẦU", "Slide 11–13", "Đối chuẩn thiết kế · Ma trận truy vết Evidence → Requirement · Bộ 12 CORE + 8 EXT + NFRs.", C_ORANGE),
        ("04 — MÔ HÌNH HÓA", "Slide 14–16", "Actors & SoD · Quy trình kép (Content & Learner) · Mô hình 15 Canonical Use Cases chuẩn.", C_GREEN),
        ("05 — HIỆN TRẠNG & KẾT LUẬN", "Slide 17", "Kết quả đạt được, tiến độ các Milestone kỹ thuật (Phase 0, M1, M2) và giới hạn nghiên cứu.", C_RED),
    ]
    
    for i, (p_title, p_slide, p_desc, p_col) in enumerate(parts):
        top_pos = Inches(1.5 + i * 1.05)
        add_card(s2, Inches(1.8), top_pos, Inches(10.7), Inches(0.95), bg_color=C_WHITE, border_color=C_BORDER)
        
        # Left tag
        tag = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.8), top_pos, Inches(0.15), Inches(0.95))
        tag.fill.solid()
        tag.fill.fore_color.rgb = p_col
        tag.line.fill.background()
        
        tbox = s2.shapes.add_textbox(Inches(2.1), top_pos + Inches(0.08), Inches(10.2), Inches(0.8))
        tf = tbox.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{p_title}  ({p_slide})"
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = p_col
        
        p2 = tf.add_paragraph()
        p2.text = p_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_DARK
        
    add_footer(s2, "Người học → Stakeholders → Nghiên cứu bổ trợ → Yêu cầu → Mô hình hệ thống")

    # Save initial
    prs.save('docs/scratch/deck_draft.pptx')
    print('Slide 1-2 built successfully!')

create_deck()
