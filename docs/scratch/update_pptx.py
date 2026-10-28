import shutil
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

pptx_path = r"docs/diagrams/lingora-modeling/deck/Lingora-System-Modeling-Presentation-v13.pptx"

prs = Presentation(pptx_path)

updates = [
    {
        'slide_index': 6, # Slide 7
        'img_path': r"docs/diagrams/lingora-modeling/out/05b-activity-placement-branches.png",
        'max_w': 11.733,
        'max_h': 4.900,
        'box_left': 0.800,
        'box_top': 1.250
    },
    {
        'slide_index': 10, # Slide 11
        'img_path': r"docs/diagrams/lingora-modeling/out/07b-activity-jump-result.png",
        'max_w': 11.733,
        'max_h': 4.900,
        'box_left': 0.800,
        'box_top': 1.250
    },
    {
        'slide_index': 13, # Slide 14
        'img_path': r"docs/diagrams/lingora-modeling/out/09-class-content.png",
        'max_w': 11.733,
        'max_h': 4.800,
        'box_left': 0.800,
        'box_top': 1.250
    },
    {
        'slide_index': 14, # Slide 15
        'img_path': r"docs/diagrams/lingora-modeling/out/10-class-learner-assessment.png",
        'max_w': 11.733,
        'max_h': 4.800,
        'box_left': 0.800,
        'box_top': 1.250
    }
]

for item in updates:
    slide = prs.slides[item['slide_index']]
    # Find existing picture shape
    pic_shape = None
    for s in slide.shapes:
        if s.shape_type == 13: # Picture
            pic_shape = s
            break
            
    if pic_shape is not None:
        # Delete old picture shape
        sp = pic_shape._element
        sp.getparent().remove(sp)
    
    # Calculate exact placement
    im = Image.open(item['img_path'])
    w_px, h_px = im.size
    orig_ratio = w_px / h_px
    
    max_w = item['max_w']
    max_h = item['max_h']
    
    h_fit = max_h
    w_fit = h_fit * orig_ratio
    if w_fit > max_w:
        w_fit = max_w
        h_fit = w_fit / orig_ratio
        
    left = item['box_left'] + (max_w - w_fit) / 2
    top = item['box_top'] + (max_h - h_fit) / 2
    
    # Add new picture
    new_pic = slide.shapes.add_picture(
        item['img_path'],
        Inches(left),
        Inches(top),
        width=Inches(w_fit),
        height=Inches(h_fit)
    )
    print(f"Slide {item['slide_index']+1}: replaced with {item['img_path']}, placed w={w_fit:.3f}, h={h_fit:.3f}, ratio={w_fit/h_fit:.4f}")

prs.save(pptx_path)
print("Saved presentation to:", pptx_path)

# Also copy to Downloads
dest_downloads = r"C:\Users\PC\Downloads\Lingora-System-Modeling-Presentation-v13.pptx"
shutil.copy2(pptx_path, dest_downloads)
print("Copied updated presentation to:", dest_downloads)
