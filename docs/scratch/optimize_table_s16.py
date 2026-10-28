import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

pptx_path = r"docs/diagrams/lingora-modeling/deck/Lingora-System-Modeling-Presentation-v13.pptx"
prs = Presentation(pptx_path)

slide_16 = prs.slides[15]
table_shape = slide_16.shapes[2]
table = table_shape.table

# Optimized widths summing to 11.733 inches
table.columns[0].width = Inches(0.85)  # UC ID
table.columns[1].width = Inches(3.70)  # Tên Ca Sử Dụng Chuẩn Tắc
table.columns[2].width = Inches(2.25)  # Sơ đồ hoạt động (slide)
table.columns[3].width = Inches(2.80)  # Lớp Dữ Liệu
table.columns[4].width = Inches(2.133) # Phạm Vi Triển Khai

# Text refinements to eliminate awkward multi-line wraps
text_adjustments = {
    4: (1, 'Xem lộ trình học & thi nhảy cấp'),
    5: (1, 'Học bài, làm bài tập & làm lại có kiểm soát'),
    9: (1, 'Kiểm duyệt chất lượng nội dung bài học'),
    10: (1, 'Xuất bản bài học đã phê duyệt'),
}

for row_idx, (col_idx, new_text) in text_adjustments.items():
    table.rows[row_idx].cells[col_idx].text = new_text

for i, row in enumerate(table.rows):
    for j, cell in enumerate(row.cells):
        cell.margin_top = Inches(0.01)
        cell.margin_bottom = Inches(0.01)
        cell.margin_left = Inches(0.04)
        cell.margin_right = Inches(0.04)
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.name = "Calibri"
                if i == 0:
                    run.font.size = Pt(12)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                elif j == 0:
                    run.font.size = Pt(12)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(37, 99, 235)
                elif j == 4:
                    run.font.size = Pt(12)
                    run.font.bold = True
                    if "MVP" in run.text:
                        run.font.color.rgb = RGBColor(21, 128, 61)
                    else:
                        run.font.color.rgb = RGBColor(100, 116, 139)
                else:
                    run.font.size = Pt(12)
                    run.font.bold = False
                    run.font.color.rgb = RGBColor(51, 65, 85)

# Move footnote to safe position
fn_shape = slide_16.shapes[3]
fn_shape.top = Inches(6.85)

prs.save(pptx_path)
print("Saved optimized Slide 16 table.")

dest_downloads = r"C:\Users\PC\Downloads\Lingora-System-Modeling-Presentation-v13.pptx"
shutil.copy2(pptx_path, dest_downloads)
print("Copied to Downloads.")
