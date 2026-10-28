import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

pptx_path = 'docs/diagrams/lingora-modeling/deck/Lingora-System-Modeling-Presentation-v13.pptx'
prs = Presentation(pptx_path)

slides_to_check = [0, 5, 9, 10, 11, 12, 13, 14, 15] # 0-indexed: 1, 6, 10, 11, 12, 13, 14, 15, 16

for idx in slides_to_check:
    sl = prs.slides[idx]
    print(f"=== SLIDE {idx+1} ===")
    for j, s in enumerate(sl.shapes):
        if s.has_text_frame:
            txt = s.text_frame.text.strip().replace('\n', ' \\n ')
            if txt:
                print(f"  Shape {j} ({s.name}): {txt}")
    if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
        print(f"  NOTES: {sl.notes_slide.notes_text_frame.text.strip()}")
