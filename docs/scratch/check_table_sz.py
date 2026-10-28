import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation('docs/diagrams/lingora-modeling/deck/Lingora-System-Modeling-Presentation-v13.pptx')
table = prs.slides[15].shapes[2].table

missing_sz = []
for i, row in enumerate(table.rows):
    for j, cell in enumerate(row.cells):
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                rPr = run._r.rPr
                sz = rPr.get('sz') if rPr is not None else None
                if sz is None:
                    missing_sz.append((i, j, run.text))

print(f"Total runs without explicit sz in table: {len(missing_sz)}")
for item in missing_sz:
    print(item)
