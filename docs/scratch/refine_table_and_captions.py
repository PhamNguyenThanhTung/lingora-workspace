import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

pptx_path = r"docs/diagrams/lingora-modeling/deck/Lingora-System-Modeling-Presentation-v13.pptx"
prs = Presentation(pptx_path)

# Fix Slide 16 table font sizes and column widths
slide_16 = prs.slides[15]
table_shape = slide_16.shapes[2]
table = table_shape.table

# Adjust column widths so Vietnamese names don't wrap to 2 lines
# Total width = 11.733 inches
# Old: [0.95, 3.5, 2.3, 2.6, 2.383]
# New: [0.90, 4.35, 1.95, 2.35, 2.183] -> Sum = 11.733
table.columns[0].width = Inches(0.90)
table.columns[1].width = Inches(4.35)
table.columns[2].width = Inches(1.95)
table.columns[3].width = Inches(2.35)
table.columns[4].width = Inches(2.183)

for i, row in enumerate(table.rows):
    for j, cell in enumerate(row.cells):
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        cell.margin_left = Inches(0.05)
        cell.margin_right = Inches(0.05)
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.name = "Calibri"
                if i == 0:
                    run.font.size = Pt(12)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                elif j == 0:
                    run.font.size = Pt(12)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(37, 99, 235) # Blue
                elif j == 4:
                    run.font.size = Pt(12)
                    run.font.bold = True
                    if "MVP" in run.text:
                        run.font.color.rgb = RGBColor(21, 128, 61) # Green
                    else:
                        run.font.color.rgb = RGBColor(100, 116, 139) # Gray
                else:
                    run.font.size = Pt(12)
                    run.font.bold = False
                    run.font.color.rgb = RGBColor(51, 65, 85) # Slate

# Check footnote position on Slide 16
fn_shape = slide_16.shapes[3]
fn_shape.top = Inches(6.85)

# Fix Slide 12 and 13 caption font sizes
for s_idx in [11, 12]:
    sl = prs.slides[s_idx]
    caption_shape = sl.shapes[4]
    for p in caption_shape.text_frame.paragraphs:
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(12.5)

prs.save(pptx_path)
print("Saved slide 16 and caption refinements.")

dest_downloads = r"C:\Users\PC\Downloads\Lingora-System-Modeling-Presentation-v13.pptx"
shutil.copy2(pptx_path, dest_downloads)
print("Copied to Downloads.")
