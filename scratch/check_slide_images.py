import sys, hashlib, os
from pptx import Presentation

prs = Presentation('docs/diagrams/lingora-modeling/deck/Lingora-System-Modeling-Presentation-v13.pptx')
out_dir = 'docs/diagrams/lingora-modeling/out'
out_hashes = {}
for f in os.listdir(out_dir):
    if f.endswith('.png'):
        path = os.path.join(out_dir, f)
        h = hashlib.md5(open(path, 'rb').read()).hexdigest()
        out_hashes[h] = f

for i, sl in enumerate(prs.slides):
    for shape in sl.shapes:
        if shape.shape_type == 13:
            img_hash = hashlib.md5(shape.image.blob).hexdigest()
            print(f"Slide {i+1} ({shape.name}): {out_hashes.get(img_hash, 'UNKNOWN')}")
