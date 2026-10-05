import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

pptx_path = r"docs/diagrams/lingora-modeling/deck/Lingora-System-Modeling-Presentation-v13.pptx"
prs = Presentation(pptx_path)

def set_run_style(run, size_pt, bold=False, color_rgb=(51, 65, 85)):
    run.font.name = "Calibri"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)

# Slide 9: Shape 9
s9_sh9 = prs.slides[8].shapes[9]
for p in s9_sh9.text_frame.paragraphs:
    for r in p.runs:
        set_run_style(r, 12, False, (51, 65, 85))

# Slide 10: Shape 7, 9, 11
s10 = prs.slides[9]
for sh_idx in [7, 9, 11]:
    sh = s10.shapes[sh_idx]
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            set_run_style(r, 12, False, (51, 65, 85))

# Slide 11: Shape 4
s11_sh4 = prs.slides[10].shapes[4]
for p in s11_sh4.text_frame.paragraphs:
    for r in p.runs:
        set_run_style(r, 12, False, (51, 65, 85))

# Slide 12: Shape 4
s12_sh4 = prs.slides[11].shapes[4]
for p in s12_sh4.text_frame.paragraphs:
    for r in p.runs:
        set_run_style(r, 12, False, (51, 65, 85))

# Slide 13: Shape 4
s13_sh4 = prs.slides[12].shapes[4]
for p in s13_sh4.text_frame.paragraphs:
    for r in p.runs:
        set_run_style(r, 12, False, (51, 65, 85))

# Slide 14: Shape 4
s14_sh4 = prs.slides[13].shapes[4]
for p in s14_sh4.text_frame.paragraphs:
    for r in p.runs:
        set_run_style(r, 12, False, (51, 65, 85))

# Slide 15: Shape 4
s15_sh4 = prs.slides[14].shapes[4]
for p in s15_sh4.text_frame.paragraphs:
    for r in p.runs:
        set_run_style(r, 12, False, (51, 65, 85))

prs.save(pptx_path)
print("Saved explicit 12pt styles to all modified text shapes.")

dest_downloads = r"C:\Users\PC\Downloads\Lingora-System-Modeling-Presentation-v13.pptx"
shutil.copy2(pptx_path, dest_downloads)
print("Copied to Downloads.")
