import shutil
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

pptx_path = r"docs/diagrams/lingora-modeling/deck/Lingora-System-Modeling-Presentation-v13.pptx"
prs = Presentation(pptx_path)

# 1. Update Slide Images
image_updates = [
    {
        'slide_index': 5, # Slide 6: 05a-activity-placement-flow
        'img_path': r"docs/diagrams/lingora-modeling/out/05a-activity-placement-flow.png",
        'max_w': 8.500,
        'max_h': 4.900,
        'box_left': 0.800,
        'box_top': 1.250
    },
    {
        'slide_index': 10, # Slide 11: 07b-activity-jump-result
        'img_path': r"docs/diagrams/lingora-modeling/out/07b-activity-jump-result.png",
        'max_w': 11.733,
        'max_h': 4.900,
        'box_left': 0.800,
        'box_top': 1.250
    },
    {
        'slide_index': 11, # Slide 12: 08a-activity-governance-main
        'img_path': r"docs/diagrams/lingora-modeling/out/08a-activity-governance-main.png",
        'max_w': 11.733,
        'max_h': 4.900,
        'box_left': 0.800,
        'box_top': 1.250
    },
    {
        'slide_index': 12, # Slide 13: 08b-activity-governance-branches
        'img_path': r"docs/diagrams/lingora-modeling/out/08b-activity-governance-branches.png",
        'max_w': 11.733,
        'max_h': 4.900,
        'box_left': 0.800,
        'box_top': 1.250
    },
    {
        'slide_index': 14, # Slide 15: 10-class-learner-assessment
        'img_path': r"docs/diagrams/lingora-modeling/out/10-class-learner-assessment.png",
        'max_w': 11.733,
        'max_h': 4.800,
        'box_left': 0.800,
        'box_top': 1.250
    }
]

for item in image_updates:
    slide = prs.slides[item['slide_index']]
    pic_shape = None
    for s in slide.shapes:
        if s.shape_type == 13: # Picture
            pic_shape = s
            break
            
    if pic_shape is not None:
        sp = pic_shape._element
        sp.getparent().remove(sp)
    
    im = Image.open(item['img_path'])
    w_px, h_px = im.size
    orig_ratio = w_px / h_px
    
    max_w = item['max_w']
    max_h = item['max_h']
    
    h_fit = max_h
    w_fit = h_fit * orig_ratio
    if w_fit > max_w:
        w_fit = max_w
        h_fit = w_fit / orig_ratio
        
    left = item['box_left'] + (max_w - w_fit) / 2
    top = item['box_top'] + (max_h - h_fit) / 2
    
    new_pic = slide.shapes.add_picture(
        item['img_path'],
        Inches(left),
        Inches(top),
        width=Inches(w_fit),
        height=Inches(h_fit)
    )
    print(f"Slide {item['slide_index']+1}: replaced image {item['img_path']}, placed w={w_fit:.3f}, h={h_fit:.3f}, ratio={w_fit/h_fit:.4f}")

# 2. Update Slide Texts & Notes

# Slide 1 (Index 0): Notes
slide_1 = prs.slides[0]
if slide_1.has_notes_slide:
    slide_1.notes_slide.notes_text_frame.text = (
        "Em chào thầy cô và các bạn. Hôm nay nhóm em trình bày phần mô hình hóa hệ thống của Lingora, "
        "nền tảng học ngoại ngữ trực tuyến, bản thí điểm cho tiếng Anh trình độ A1. "
        "Phần này đi theo bốn bước: use case, sơ đồ hoạt động, sơ đồ lớp, và cuối cùng là bảng truy vết "
        "để truy vết use case sang sơ đồ và dữ liệu."
    )
print("Updated Slide 1 notes.")

# Slide 6 (Index 5): Notes
slide_6 = prs.slides[5]
if slide_6.has_notes_slide:
    slide_6.notes_slide.notes_text_frame.text = (
        "Sơ đồ hoạt động đầu tiên là bài xếp lớp đầu vào. Học viên làm đề 30 câu, hệ thống tính điểm thô "
        "từ 0 đến 30 rồi áp dụng Placement Mapping quy đổi ra cấp bắt đầu, Course và Unit bắt đầu. "
        "Ở bản thí điểm nhóm em chỉ xếp trong A1, điểm quy đổi ra Unit 1 đến 6; việc xếp lên các cấp cao hơn "
        "chưa có dữ liệu kiểm định. Cuối cùng hệ thống lưu kết quả xếp lớp cố định và cập nhật hồ sơ năng lực "
        "cùng vị trí bắt đầu vào Profile. Hai phần cuối này là mục tiêu của giai đoạn 2."
    )
print("Updated Slide 6 notes.")

# Slide 9 (Index 8): Body text & Notes
slide_9 = prs.slides[8]
for s in slide_9.shapes:
    if s.has_text_frame and "Tối đa max(1, floor(N × 0,10))" in s.text_frame.text:
        s.text_frame.text = "Tối đa max(1, floor(N × 0,10)) lượt (với N là số câu hỏi trong bài). Hết lượt sẽ báo 'No quota' và câu đó được 0 điểm."
if slide_9.has_notes_slide:
    slide_9.notes_slide.notes_text_frame.text = (
        "Khi học viên trả lời sai, hệ thống hiện lời giải thích đáp án trước. "
        "Sau đó mới cho làm lại, với số lượt tối đa là lớn hơn giữa 1 và 10 phần trăm số câu trong bài học. "
        "Ví dụ bài 10 câu thì được làm lại 1 câu. Hết lượt thì báo hết quota và câu đó được 0 điểm. "
        "Nếu làm lại đúng, điểm chỉ được tính tối đa 70 phần trăm. Nhóm em đặt như vậy để người làm đúng ngay lần đầu vẫn có lợi thế."
    )
print("Updated Slide 9 text and notes.")

# Slide 10 (Index 9): Body text & Notes
slide_10 = prs.slides[9]
for s in slide_10.shapes:
    if s.has_text_frame and "Nếu đã thi trượt Jump trong 24 giờ qua" in s.text_frame.text:
        s.text_frame.text = "Khóa 24 giờ trên toàn tài khoản Learner: nếu đã thi trượt Jump trong 24 giờ qua thì chặn làm bài và hiện đồng hồ đếm ngược."
    elif s.has_text_frame and "Nhảy 1 Unit: 20 câu. 2–3 Unit: 25–30 câu." in s.text_frame.text:
        s.text_frame.text = "Quy mô đề Jump: nhảy 1 Unit gồm 20 câu; nhảy 2 Unit gồm 25 câu; nhảy 3 Unit gồm 30 câu; trên 3 Unit gồm 40 câu."
if slide_10.has_notes_slide:
    slide_10.notes_slide.notes_text_frame.text = (
        "Thi nhảy cấp chỉ bỏ qua Unit trong cùng một Course, không thay đổi cấp CEFR. "
        "Bước đầu tiên là kiểm tra cooldown: nếu học viên vừa thi trượt Jump trong 24 giờ qua, "
        "hệ thống khóa tính năng Jump trên toàn tài khoản và hiện đồng hồ đếm ngược. "
        "Nếu hợp lệ, số câu tăng theo số Unit muốn nhảy: 20 câu khi nhảy 1 Unit, 25 câu khi nhảy 2 Unit, "
        "30 câu khi nhảy 3 Unit, và 40 câu khi nhảy trên 3 Unit. Bài được chấm tự động và lưu kết quả."
    )
print("Updated Slide 10 text and notes.")

# Slide 11 (Index 10): Caption & Notes
slide_11 = prs.slides[10]
for s in slide_11.shapes:
    if s.has_text_frame and "Đỗ (từ 95%):" in s.text_frame.text:
        s.text_frame.text = (
            "Đỗ (từ 95%): gán SKIPPED_BY_JUMP (by_jump = true) cho phần bỏ qua, mở Unit mục tiêu trong cùng Course. CEFR giữ nguyên.\n"
            "Trượt (dưới 95%): kích hoạt cooldown 24 giờ trên tài khoản Learner (khóa tính năng Jump 24h), giữ vị trí học hiện tại và học tiếp theo lộ trình. CEFR giữ nguyên."
        )
if slide_11.has_notes_slide:
    slide_11.notes_slide.notes_text_frame.text = (
        "Sau khi chấm có hai nhánh. Từ 95 phần trăm trở lên là đỗ: phần được bỏ qua gắn nhãn vượt bằng Jump "
        "và Unit mục tiêu mở ngay, trong cùng Course. Dưới 95 phần trăm là trượt: tài khoản học viên bị khóa tính năng Jump 24 giờ, "
        "vị trí học giữ nguyên và học viên tiếp tục theo lộ trình. Ở cả hai nhánh, cấp CEFR đều giữ nguyên."
    )
print("Updated Slide 11 caption and notes.")

# Slide 12 (Index 11): Caption & Notes
slide_12 = prs.slides[11]
for s in slide_12.shapes:
    if s.has_text_frame and "Moderator chuyển bài sang APPROVED" in s.text_frame.text:
        s.text_frame.text = (
            "Org Manager phân công soạn bài. Creator nộp duyệt; nếu bị yêu cầu sửa (CHANGES_REQUESTED) thì chỉnh sửa và gửi duyệt lại cho Moderator.\n"
            "Chỉ Org Manager được xuất bản (Publish) bài đã APPROVED khi Unit thỏa mãn điều kiện đủ 4/4 bài."
        )
if slide_12.has_notes_slide:
    slide_12.notes_slide.notes_text_frame.text = (
        "Đây là luồng duyệt và xuất bản bài học, có bốn làn: Org Manager, Creator, Moderator và hệ thống. "
        "Org Manager giao bài và phân công. Creator soạn xong thì gửi kiểm duyệt. Moderator thẩm định chất lượng: "
        "nếu đạt thì phê duyệt sang APPROVED; nếu chưa đạt thì yêu cầu chỉnh sửa (CHANGES_REQUESTED) kèm góp ý chi tiết, "
        "Creator sửa bài và gửi nộp duyệt lại tạo thành vòng lặp. Khi bài đã APPROVED, Org Manager kiểm tra điều kiện Unit "
        "đủ 4 trên 4 bài rồi thực hiện xuất bản, sau đó hệ thống ghi nhận nhật ký AuditLog."
    )
print("Updated Slide 12 caption and notes.")

# Slide 13 (Index 12): Caption & Notes
slide_13 = prs.slides[12]
for s in slide_13.shapes:
    if s.has_text_frame and "Clone-to-Draft v2:" in s.text_frame.text:
        s.text_frame.text = (
            "Clone-to-Draft v2: sửa bài đã PUBLISHED sẽ tạo bản nháp v2. v1 vẫn phục vụ học viên cho đến khi v2 được Moderator duyệt APPROVED và Org Manager xuất bản mới (Archive v1).\n"
            "Xuất bản khẩn cấp: Administrator chỉ can thiệp xuất bản bài đã APPROVED kèm lý do giải trình bắt buộc và lưu vết AuditLog."
        )
if slide_13.has_notes_slide:
    slide_13.notes_slide.notes_text_frame.text = (
        "Còn nếu cần sửa một bài đã xuất bản thì sao? Nhóm em dùng cơ chế Clone-to-Draft. "
        "Hệ thống tạo bản nháp v2 từ bản đang chạy, còn bản v1 vẫn phục vụ học viên trong lúc sửa. "
        "Bản v2 sau khi sửa vẫn phải qua bước kiểm duyệt của Moderator để đạt APPROVED. "
        "Sau đó, Org Manager xuất bản v2 và hệ thống tự động lưu trữ v1. "
        "Nhánh thay thế khẩn cấp do Administrator thực hiện cũng chỉ áp dụng cho bài đã APPROVED, "
        "bắt buộc có lý do giải trình và luôn được ghi nhật ký AuditLog."
    )
print("Updated Slide 13 caption and notes.")

# Slide 14 (Index 13): Caption & Notes
slide_14 = prs.slides[13]
for s in slide_14.shapes:
    if s.has_text_frame and "Chú giải: CURRENT = nét liền" in s.text_frame.text:
        s.text_frame.text = (
            "Chú giải: CURRENT = nét liền (đã có trong schema.prisma)  ·  <<TARGET>> = nét đứt (dự kiến Phase 2)\n"
            "Course.cefr dùng enum Cefr (A1–C2). Mỗi Course có 1..N Unit; bản thí điểm English A1 có 6 Unit, không áp cho cấp khác.\n"
            "Lesson.is_published (CURRENT) phục vụ hiển thị nhanh; LessonRevision (TARGET) quản lý 6 trạng thái vòng đời. Unit chỉ hiện khi đủ 4/4 bài Published."
        )
if slide_14.has_notes_slide:
    slide_14.notes_slide.notes_text_frame.text = (
        "Sang sơ đồ lớp nội dung. Nét liền là phần đã có trong schema, nét đứt là mục tiêu giai đoạn 2. "
        "Mỗi Course gắn với một cấp CEFR và có từ 1 đến N Unit; bản thí điểm A1 có 6 Unit, không áp cho cấp khác. "
        "Course chứa Unit, Unit chứa Lesson, Lesson có Exercise và từ vựng. "
        "Ở mô hình hiện tại, Lesson.is_published là cờ boolean phục vụ hiển thị nhanh cho học viên, "
        "còn ở Phase 2 lớp LessonRevision sẽ quản lý đầy đủ 6 trạng thái kiểm duyệt và xuất bản. "
        "Unit chỉ hiển thị khi có đủ 4 trên 4 bài học ở trạng thái Published."
    )
print("Updated Slide 14 caption and notes.")

# Slide 15 (Index 14): Caption & Notes
slide_15 = prs.slides[14]
for s in slide_15.shapes:
    if s.has_text_frame and "Chú giải: CURRENT = nét liền" in s.text_frame.text:
        s.text_frame.text = (
            "Chú giải: CURRENT = nét liền (đã có trong schema.prisma)  ·  <<TARGET>> = nét đứt (dự kiến Phase 2)\n"
            "UserProgress (tiến độ bài học, retry_count) tách khỏi SkillProfile (năng lực verified_cefr từ AssessmentResult). Profile quản lý jump_cooldown_until.\n"
            "SupportTicket (TARGET) phục vụ điều phối hỗ trợ (UC-14). AssessmentResult gồm: PLACEMENT, JUMP, LEVEL_ADVANCEMENT."
        )
if slide_15.has_notes_slide:
    slide_15.notes_slide.notes_text_frame.text = (
        "Phía học viên, nhóm em tách UserProgress, tức tiến độ học bài, khỏi SkillProfile, tức năng lực đã kiểm chứng. "
        "Quota làm lại trong bài được theo dõi qua trường retry_count trong UserProgress, còn thời gian cooldown nhảy cấp 24 giờ "
        "được lưu tại trường jump_cooldown_until trong Profile của người học. "
        "Kết quả đánh giá lưu ở AssessmentResult với ba mục đích: xếp lớp, nhảy cấp và chuyển cấp. "
        "Lớp SupportTicket được bổ sung ở mức TARGET để thể hiện việc tiếp nhận và xử lý phiếu hỗ trợ kỹ thuật của UC-14."
    )
print("Updated Slide 15 caption and notes.")

# Slide 16 (Index 15): Table columns translation & Notes
slide_16 = prs.slides[15]
# Mapping UC Vietnamese titles
uc_translations = {
    'UC-01': 'Đăng nhập và phân quyền phiên làm việc',
    'UC-02': 'Đặt mục tiêu học tập và thời gian học',
    'UC-03': 'Làm bài xếp lớp và xác định xuất phát điểm',
    'UC-04': 'Xem lộ trình học tập và thi nhảy cấp',
    'UC-05': 'Học bài, làm bài tập và làm lại có kiểm soát',
    'UC-06': 'Ôn tập từ vựng ngắt quãng (SRS)',
    'UC-07': 'Theo dõi tiến độ, mức độ thuần thục và streak',
    'UC-08': 'Soạn thảo bản nháp bài học và bài tập',
    'UC-09': 'Thẩm định chất lượng và kiểm duyệt nội dung',
    'UC-10': 'Xuất bản bài học đã được phê duyệt',
    'UC-11': 'Quản lý tài khoản người dùng và 6 vai trò',
    'UC-12': 'Xem báo cáo thống kê hoạt động hệ thống',
    'UC-13': 'Tra cứu và kiểm toán nhật ký an ninh',
    'UC-14': 'Tiếp nhận và xử lý phiếu hỗ trợ kỹ thuật',
    'UC-15': 'Quản lý nhóm học viên và tiến độ nhóm (mở rộng)'
}

for s in slide_16.shapes:
    if s.has_table:
        table = s.table
        for row in table.rows:
            uc_id = row.cells[0].text.strip()
            if uc_id in uc_translations:
                row.cells[1].text = uc_translations[uc_id]
                # If UC-14, update data model cell
                if uc_id == 'UC-14':
                    row.cells[3].text = 'Profile, SupportTicket'

if slide_16.has_notes_slide:
    slide_16.notes_slide.notes_text_frame.text = (
        "Bảng này nối ba thứ lại với nhau: từng ca sử dụng chuẩn tắc bằng tiếng Việt tương ứng với slide nào "
        "và sử dụng những thực thể dữ liệu nào. Ví dụ UC-04 ứng với slide 10 và 11, truy vết sang Course, Unit, Lesson và UserProgress. "
        "Dấu gạch nghĩa là use case đó không có sơ đồ hoạt động riêng. "
        "UC-14 xử lý phiếu hỗ trợ được ánh xạ tới Profile và thực thể dữ liệu SupportTicket, "
        "còn UC-15 là hướng mở rộng cho giai đoạn sau."
    )
print("Updated Slide 16 table and notes.")

# Slide 17 (Index 16): Notes update
slide_17 = prs.slides[16]
if slide_17.has_notes_slide:
    slide_17.notes_slide.notes_text_frame.text = (
        "Để kết lại: thứ nhất, cấp CEFR, Course, Unit và tiến độ là bốn khái niệm tách bạch, "
        "và 15 use case, 8 sơ đồ hoạt động, 2 sơ đồ lớp khớp nhau hoàn toàn qua bảng truy vết. "
        "Thứ hai, xếp lớp chỉ quyết định Unit bắt đầu, nhảy cấp chỉ bỏ qua Unit với cooldown 24 giờ trên tài khoản, "
        "còn chuyển cấp là hướng mở rộng chưa chốt. "
        "Thứ ba, quản trị tách quyền rõ ràng, sửa bài đã xuất bản dùng Clone-to-Draft v2 luôn qua bước kiểm duyệt "
        "và xuất bản khẩn cấp luôn có nhật ký AuditLog. "
        "Em xin cảm ơn thầy cô và các bạn, nhóm em sẵn sàng nhận câu hỏi."
    )
print("Updated Slide 17 notes.")

prs.save(pptx_path)
print("Saved presentation successfully:", pptx_path)

dest_downloads = r"C:\Users\PC\Downloads\Lingora-System-Modeling-Presentation-v13.pptx"
shutil.copy2(pptx_path, dest_downloads)
print("Copied updated presentation to:", dest_downloads)
