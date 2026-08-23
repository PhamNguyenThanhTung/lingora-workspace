from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\btl")
SOURCE = ROOT / "docs" / "report" / "lingora-project-1-lecturer-format.md"
OUTPUT = ROOT / "docs" / "report" / "Lingora-Project-1-Lecturer-Format-Final.docx"


def set_run_font(run, size=13, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for node in (fld_begin, instr, fld_sep, text, fld_end):
        run._r.append(node)
    return run


def add_inline(paragraph, text, size=13):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        italic = (part.startswith("*") and part.endswith("*")) and not bold
        clean = part[2:-2] if bold else part[1:-1] if (code or italic) else part
        run = paragraph.add_run(clean)
        set_run_font(run, size=size, bold=bold, italic=italic)
        if code:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(max(size - 1, 9))


def style_body(paragraph, indent=True):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.3
    fmt.space_after = Pt(6)
    fmt.space_before = Pt(0)
    if indent:
        fmt.first_line_indent = Cm(1.0)
    fmt.widow_control = True


def add_cover(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    for _ in range(2):
        doc.add_paragraph()
    for text, size in [
        ("TRƯỜNG ĐẠI HỌC [BỔ SUNG TÊN TRƯỜNG]", 14),
        ("KHOA [BỔ SUNG TÊN KHOA]", 13),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=size, bold=True)
        p.paragraph_format.space_after = Pt(4)
    doc.add_paragraph()
    for text, size in [
        ("BÁO CÁO PROJECT 1", 18),
        ("MÔN PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", 15),
        ("ĐỀ TÀI", 14),
        ("PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG LINGORA", 17),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=size, bold=True)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
    doc.add_paragraph()
    meta = [
        "Giảng viên hướng dẫn: [Bổ sung]",
        "Sinh viên thực hiện: Phạm Nguyễn Thanh Tùng và các thành viên nhóm",
    ]
    for text in meta:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(3.5)
        r = p.add_run(text)
        set_run_font(r, size=13)
        p.paragraph_format.space_after = Pt(6)
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HÀ NỘI - 2026")
    set_run_font(r, size=13, bold=True)
    doc.add_page_break()


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MỤC LỤC")
    set_run_font(r, size=16, bold=True)
    p.paragraph_format.space_after = Pt(12)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Nhấn Ctrl+A, F9 trong Word để cập nhật mục lục")
    for run in toc.runs:
        set_run_font(run, size=12)
    doc.add_page_break()


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.27)
    section.footer_distance = Cm(1.27)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in [
        ("Heading 1", 15, 12, 6),
        ("Heading 2", 14, 10, 5),
        ("Heading 3", 13, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = header.add_run("BÁO CÁO PROJECT 1 - HỆ THỐNG LINGORA")
    set_run_font(hr, size=10, italic=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Trang ")
    set_run_font(fr, size=10)
    page_run = add_field(footer, "PAGE", "1")
    set_run_font(page_run, size=10)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    usable = Cm(15.92)
    width = usable / cols
    for ridx, row in enumerate(rows):
        for cidx in range(cols):
            cell = table.cell(ridx, cidx)
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ridx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            add_inline(p, row[cidx] if cidx < len(row) else "", size=11)
            if ridx == 0:
                set_cell_shading(cell, "D9E2F3")
                for run in p.runs:
                    run.bold = True
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_markdown_content(doc, text):
    lines = text.splitlines()
    start = 0
    for idx, line in enumerate(lines):
        if line.strip() == "# LỜI GIỚI THIỆU":
            start = idx
            break
    i = start
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if stripped.startswith("# "):
            title = stripped[2:].strip()
            if title != "LỜI GIỚI THIỆU":
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            add_inline(p, title.upper(), size=16)
            i += 1
            continue
        if stripped.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, stripped[3:].strip(), size=15)
            i += 1
            continue
        if stripped.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, stripped[4:].strip(), size=14)
            i += 1
            continue
        if stripped.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, stripped[5:].strip(), size=13)
            i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:].strip(), size=13)
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped), size=13)
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue
        if stripped.startswith("**[CHÈN HÌNH:") or stripped.startswith("`[CHÈN ẢNH:"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, stripped.strip("`*"), size=11)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(89, 89, 89)
            i += 1
            continue
        p = doc.add_paragraph()
        style_body(p)
        add_inline(p, stripped, size=13)
        i += 1


def main():
    source_text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_toc(doc)
    add_markdown_content(doc, source_text)
    core = doc.core_properties
    core.title = "Báo cáo Project 1 - Phân tích và Thiết kế Hệ thống Lingora"
    core.subject = "Requirements analysis and system design"
    core.author = "Phạm Nguyễn Thanh Tùng và nhóm Lingora"
    core.keywords = "Lingora, Project 1, requirements, system analysis"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
